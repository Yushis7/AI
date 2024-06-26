from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart,Reference

bar_value = Reference(ws,min_row=2,max_row=11,min_col=2,max_col=3)
bar_chart = BarChart() #차트 종류 bar,line,pie...
bar_chart.add_data(bar_value) #차트 데이터 추가

ws.add_chart(bar_chart,"E1") #차트 넣을 위치 정의

wb.save("sample_chart.xlsx")
